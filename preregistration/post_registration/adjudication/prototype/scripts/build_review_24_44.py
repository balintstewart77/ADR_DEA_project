"""Build the blind review sheet for splitting earlier mechanism-24 findings (ADJ-091).

Mechanism 24 ("label assigned from an incidental mention or background data")
was split on 2026-09-28: 44 now covers a label with no basis in the title or
datasets.  Findings from ADJ_0111's Stage 2 onwards used 44 directly; every
earlier finding coded 24 is reviewed once, against the rule fixed in ADJ-091,
before any tabulation by mechanism.

The sheet shows, per finding: the frozen title and datasets, the labels the
finding concerns, the Stage 1 rule conflicts it rests on (rule, labels and the
blind Stage 1 explanation; findings before the split did not record their own
labels), and the primary's notes.  The issue family and the source are
left out, and source words in the notes are masked, so the recoding cannot
lean by source.  Rows are shuffled with a fixed seed.  The key from sheet row
to record and finding is written separately and is not needed for the review.

Reads the final export of the primary pass and the harmonisation file; writes
only to ``preregistration_restricted/adjudication_formal/review_24_44/``.
Prints counts only.  Refuses to overwrite.
"""
from __future__ import annotations

import csv
import hashlib
import json
import random
import re
from datetime import datetime, timezone

from build_formal_import import OUT, ROOT
from preserve_block import _choices, flat

EXPORT = OUT / "exports" / "AdjudicationPRODUCTI_DATA_2026-09-30_0953.csv"
EXPORT_SHA256 = "4e10b3ceea2dd506743f5bae5acb0ed33ebcc7df142239e0ce3b905800b0469f"
HARMONISATION = OUT / "preservation" / "mechanism_harmonisation.csv"
DICTIONARY = ROOT / "instruments" / "adjudication_stage1_candidate.csv"
DEST = OUT / "review_24_44"
LAST_IN_SCOPE = 110            # ADJ_0110: findings from ADJ_0111's Stage 2 used 44 directly
SEED_REVIEW_24_44 = 20261016   # row order on the sheet; fixed before the sheet was built
SLOTS = range(1, 7)
SOURCE_WORDS = re.compile(r"\b(Fable(\s*5)?|GPT[-\s]?5(\.5)?|LLM|AI|(production\s+)?model'?s?|scratch[-\s]?coders?'?s?|coders?'?s?|C0[123]|humans?|majority)\b", re.I)
RULE = ("24: something in the title or datasets mentions or signals the label, but it is not central to the project. "
        "44: nothing in the title or datasets points to the label; it was inferred from outside the entry, or guessed. "
        "(27, unchanged: nothing in the entry resolves the whole layer.)")


def labels(field_rows):
    out = {}
    for r in field_rows:
        if r[0] in ("adj_f1_dom_labels", "adj_f1_purp_labels"):
            out[r[0].split("_")[2]] = dict(_choices(r[5]))
    return out


def rule_names(field_rows):
    return dict(_choices(next(r for r in field_rows if r[0] == "adj_conflict_rule_cited")[5]))


def conflict_fields(n):
    """Stage 1 field names for rule conflict n (conflict 1 carries no number)."""
    m = "" if n == 1 else str(n)
    return {"rule": f"adj_conflict_rule{m}_cited", "other": f"adj_conflict_rule{m}_other",
            "dom": f"adj_dom_conflict{m}_labels", "purp": f"adj_purp_conflict{m}_labels",
            "note": "adj_rule_conflict_note" + m}


def mask(text):
    return SOURCE_WORDS.subn("[source]", flat(text or ""))


def main():
    DEST.mkdir(parents=True, exist_ok=True)
    sheet_path, key_path = DEST / "review_sheet_24_44_v2.xlsx", DEST / "review_sheet_24_44_key_v2.csv"
    if sheet_path.exists() or key_path.exists():
        raise SystemExit("the review sheet already exists; it is built once")
    data = EXPORT.read_bytes()
    if hashlib.sha256(data).hexdigest() != EXPORT_SHA256:
        raise SystemExit("the final export does not match its recorded hash")
    rows = {r["adj_assignment_id"]: r for r in csv.DictReader(data.decode("utf-8-sig").splitlines(keepends=True))}
    field_rows = list(csv.reader(DICTIONARY.open(encoding="utf-8")))[1:]
    vocab, rules = labels(field_rows), rule_names(field_rows)
    harmonised = {(h["assignment_id"], int(h["finding"])): h["harmonised_code"]
                  for h in csv.DictReader(HARMONISATION.open(encoding="utf-8"))}

    items, later, masked = [], 0, 0
    for aid, row in sorted(rows.items()):
        n = int(aid.split("_")[1])
        for k in SLOTS:
            code = row.get(f"adj_f{k}_mech", "")
            how = "picked" if code == "24" else ("harmonised" if harmonised.get((aid, k)) == "24" else "")
            if not how:
                continue
            if n > LAST_IN_SCOPE:
                later += 1
                continue
            concerned = []
            for comp, name in (("dom", "Research Domain"), ("purp", "Analytical Purpose")):
                for c, label in vocab[comp].items():
                    if row.get(f"adj_f{k}_{comp}_labels___{c}") == "1":
                        concerned.append(f"{name}: {label}")
            notes, hits = [], 0
            for field in ("note", "comment"):
                text, n_hits = mask(row.get(f"adj_f{k}_{field}", ""))
                hits += n_hits
                notes.append(text)
            linked = []
            for n in range(1, 7):
                if row.get(f"adj_f{k}_conflicts___{n}") != "1":
                    continue
                f = conflict_fields(n)
                rule = rules.get(row.get(f["rule"], ""), "")
                if row.get(f["other"]):
                    other, n_hits = mask(row[f["other"]]); hits += n_hits
                    rule = f"{rule} ({other})" if rule else other
                conf_labels = [f"{name}: {label}" for comp, name in (("dom", "Research Domain"), ("purp", "Analytical Purpose"))
                               for c, label in vocab[comp].items() if row.get(f"{f[comp]}___{c}") == "1"]
                note, n_hits = mask(row.get(f["note"], "")); hits += n_hits
                linked.append(f"Conflict {n}. Rule: {rule or '(not recorded)'}. Labels: {'; '.join(conf_labels) or '(none)'}. "
                              f"Stage 1 explanation: {note or '(none)'}")
            masked += bool(hits)
            items.append({"assignment_id": aid, "finding": k, "recorded_as": how, "title": flat(row.get("adj_case_title", "")),
                          "datasets": flat(row.get("adj_case_datasets", "")), "labels": "; ".join(concerned) or "(none ticked)",
                          "basis_note": notes[0], "optional_note": notes[1],
                          "conflicts": "\n\n".join(linked) or "(not based on a recorded conflict)"})

    random.Random(SEED_REVIEW_24_44).shuffle(items)
    for i, it in enumerate(items, 1):
        it["review_id"] = f"R{i:03d}"

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    ws.append(["Decide each row: keep 24, or recode 44. " + RULE + " Family and source are hidden; source words in the notes read [source]."])
    ws.append([])
    header = ["review_id", "title", "datasets", "labels concerned (finding)", "linked Stage 1 conflicts (blind)", "basis note", "optional note", "decision (24 or 44)", "your comment (optional)"]
    ws.append(header)
    for it in items:
        ws.append([it["review_id"], it["title"], it["datasets"], it["labels"], it["conflicts"], it["basis_note"], it["optional_note"], "", ""])
    widths = [10, 36, 36, 26, 70, 36, 30, 14, 30]
    for col, w in zip("ABCDEFGHI", widths):
        ws.column_dimensions[col].width = w
    for r in ws.iter_rows(min_row=3):
        for cell in r:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[3]:
        cell.font = Font(bold=True)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A1:I1")
    ws.row_dimensions[1].height = 60
    ws.freeze_panes = "B4"
    dv = DataValidation(type="list", formula1='"24,44"', allow_blank=True)
    ws.add_data_validation(dv)
    if items:
        dv.add(f"H4:H{len(items) + 3}")
    wb.save(sheet_path)

    with key_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["review_id", "assignment_id", "finding", "recorded_as"], lineterminator="\n")
        w.writeheader()
        w.writerows({k: it[k] for k in ("review_id", "assignment_id", "finding", "recorded_as")} for it in items)
    digest = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
    receipt = {"built_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"), "rule": RULE,
               "scope": f"findings coded 24 (picked, or harmonised from typed text) on ADJ_0001 to ADJ_{LAST_IN_SCOPE:04d}",
               "export": EXPORT.name, "export_sha256": EXPORT_SHA256, "seed": SEED_REVIEW_24_44,
               "findings_in_scope": len(items), "records_in_scope": len({it["assignment_id"] for it in items}),
               "findings_coded_24_after_scope": later, "findings_with_masked_source_words": masked,
               "findings_with_no_label_information": sum(1 for it in items if it["labels"] == "(none ticked)" and it["conflicts"] == "(not based on a recorded conflict)"),
               "sheet_sha256": digest(sheet_path), "key_sha256": digest(key_path)}
    (DEST / "review_sheet_24_44_receipt_v2.json").write_text(json.dumps(receipt, indent=2) + "\n", encoding="utf-8")
    print(f"{len(items)} findings to review, on {receipt['records_in_scope']} records; "
          f"{later} finding(s) coded 24 after ADJ_{LAST_IN_SCOPE:04d} are out of scope; "
          f"source words masked in {masked} finding(s).")


if __name__ == "__main__":
    main()
